import argparse

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sklearn.ensemble import RandomForestRegressor
from sklearn.linear_model import LinearRegression
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.model_selection import KFold, cross_val_score
from xgboost import XGBRegressor

from preprocess import PROJECT_ROOT, prepare_training_data, split_data


EVAL_DIR = PROJECT_ROOT / "evaluation_plots"
EXCEL_PATH = PROJECT_ROOT / "model_comparison_results.xlsx"


def build_models():
    return {
        "Linear Regression": LinearRegression(),
        "Random Forest": RandomForestRegressor(
            n_estimators=200,
            max_depth=None,
            min_samples_split=2,
            min_samples_leaf=1,
            random_state=42,
            n_jobs=-1,
        ),
        "XGBoost": XGBRegressor(
            n_estimators=1000,
            learning_rate=0.01,
            max_depth=3,
            subsample=0.8,
            colsample_bytree=0.8,
            random_state=42,
            objective="reg:squarederror",
            verbosity=0,
        ),
    }


def evaluate_models(X_train, X_test, y_train, y_test, X, y):
    results = {}

    for name, model in build_models().items():
        print(f"\n{'=' * 60}")
        print(f"  {name}")
        print(f"{'=' * 60}")

        model.fit(X_train, y_train)
        y_pred = model.predict(X_test)

        y_test_orig = np.expm1(y_test)
        y_pred_orig = np.expm1(y_pred)

        results[name] = {
            "model": model,
            "y_pred": y_pred,
            "y_pred_orig": y_pred_orig,
            "RMSE (Log)": np.sqrt(mean_squared_error(y_test, y_pred)),
            "R2 Score": r2_score(y_test, y_pred),
            "MAE (Log)": mean_absolute_error(y_test, y_pred),
            "RMSE ($)": np.sqrt(mean_squared_error(y_test_orig, y_pred_orig)),
            "MAE ($)": mean_absolute_error(y_test_orig, y_pred_orig),
        }

        print(f"  RMSE (Log)  : {results[name]['RMSE (Log)']:.4f}")
        print(f"  R2 Score    : {results[name]['R2 Score']:.4f}")
        print(f"  MAE (Log)   : {results[name]['MAE (Log)']:.4f}")
        print(f"  RMSE (USD)  : ${results[name]['RMSE ($)']:,.2f}")
        print(f"  MAE (USD)   : ${results[name]['MAE ($)']:,.2f}")

    kf = KFold(n_splits=5, shuffle=True, random_state=42)
    print(f"\n{'=' * 60}")
    print("  CROSS VALIDATION (5-Fold)")
    print(f"{'=' * 60}")

    for name, model in build_models().items():
        cv_scores = cross_val_score(
            model,
            X,
            y,
            cv=kf,
            scoring="neg_mean_squared_error",
        )
        cv_rmse = np.sqrt(-cv_scores)
        results[name]["CV RMSE Mean"] = cv_rmse.mean()
        results[name]["CV RMSE Std"] = cv_rmse.std()

        print(f"\n  {name}:")
        print(f"    CV RMSE : {cv_rmse.mean():.4f} +/- {cv_rmse.std():.4f}")
        print(f"    Folds   : {[f'{s:.4f}' for s in cv_rmse]}")

    return results


def save_plots(results, y_test, feature_columns, show_plots=False):
    EVAL_DIR.mkdir(exist_ok=True)

    fig, axes = plt.subplots(1, 3, figsize=(18, 5))
    for ax, (name, info) in zip(axes, results.items()):
        ax.scatter(y_test, info["y_pred"], alpha=0.5, s=15, c="#3498db")
        min_val = min(y_test.min(), info["y_pred"].min())
        max_val = max(y_test.max(), info["y_pred"].max())
        ax.plot([min_val, max_val], [min_val, max_val], "r--", lw=2)
        ax.set_xlabel("Actual (Log Scale)")
        ax.set_ylabel("Predicted (Log Scale)")
        ax.set_title(f"{name}\nR2 = {info['R2 Score']:.4f}")
    fig.suptitle("Actual vs Predicted - All Models", fontsize=14, fontweight="bold")
    fig.tight_layout()
    fig.savefig(EVAL_DIR / "actual_vs_predicted.png", dpi=150, bbox_inches="tight")
    _finish_figure(fig, show_plots)

    fig, axes = plt.subplots(1, 3, figsize=(18, 5))
    for ax, (name, info) in zip(axes, results.items()):
        residuals = y_test - info["y_pred"]
        ax.scatter(info["y_pred"], residuals, alpha=0.5, s=15, c="#e74c3c")
        ax.axhline(y=0, color="black", linestyle="--", lw=1.5)
        ax.set_xlabel("Predicted Values")
        ax.set_ylabel("Residuals")
        ax.set_title(name)
    fig.suptitle("Residual Analysis - All Models", fontsize=14, fontweight="bold")
    fig.tight_layout()
    fig.savefig(EVAL_DIR / "residual_analysis.png", dpi=150, bbox_inches="tight")
    _finish_figure(fig, show_plots)

    fig, axes = plt.subplots(1, 2, figsize=(16, 8))
    for ax, name in zip(axes, ["Random Forest", "XGBoost"]):
        importance_df = pd.DataFrame({
            "Feature": feature_columns,
            "Importance": results[name]["model"].feature_importances_,
        }).sort_values("Importance", ascending=False)
        top15 = importance_df.head(15)
        colors = plt.cm.viridis(np.linspace(0.3, 0.9, 15))
        ax.barh(top15["Feature"][::-1], top15["Importance"][::-1], color=colors)
        ax.set_xlabel("Importance")
        ax.set_title(f"{name} - Top 15 Features")
    fig.suptitle("Feature Importance Comparison", fontsize=14, fontweight="bold")
    fig.tight_layout()
    fig.savefig(EVAL_DIR / "feature_importance.png", dpi=150, bbox_inches="tight")
    _finish_figure(fig, show_plots)

    fig, axes = plt.subplots(1, 3, figsize=(16, 6))
    model_names = list(results.keys())
    colors = ["#2ecc71", "#3498db", "#e67e22"]

    metric_specs = [
        ("R2 Score", "R2 Score (higher = better)", False),
        ("RMSE (Log)", "RMSE - Log Scale (lower = better)", False),
        ("CV RMSE Mean", "Cross-Validation RMSE (lower = better)", True),
    ]

    for ax, (metric, title, use_error), in zip(axes, metric_specs):
        values = [results[name][metric] for name in model_names]
        errors = [results[name]["CV RMSE Std"] for name in model_names] if use_error else None
        bars = ax.bar(model_names, values, yerr=errors, color=colors, alpha=0.85, capsize=5)
        for bar, val in zip(bars, values):
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_height() + 0.002,
                f"{val:.4f}",
                ha="center",
                va="bottom",
                fontweight="bold",
            )
        ax.set_title(title)
        ax.set_ylim(min(values) - 0.05, max(values) + 0.05)

    fig.suptitle("Model Performance Comparison", fontsize=14, fontweight="bold")
    fig.tight_layout()
    fig.savefig(EVAL_DIR / "metrics_comparison.png", dpi=150, bbox_inches="tight")
    _finish_figure(fig, show_plots)


def _finish_figure(fig, show_plots):
    if show_plots:
        plt.show()
    else:
        plt.close(fig)


def export_results(results, X_train, X_test, y_test, feature_columns):
    wb = Workbook()

    header_font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    best_fill = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    center_align = Alignment(horizontal="center", vertical="center")

    def style_header(ws, num_cols):
        for col_idx in range(1, num_cols + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

    def auto_width(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
                cell.border = thin_border
                cell.alignment = center_align
            ws.column_dimensions[col_letter].width = max_len + 4

    summary_data = [
        {
            "Model": name,
            "R2 Score": round(info["R2 Score"], 4),
            "RMSE (Log)": round(info["RMSE (Log)"], 4),
            "MAE (Log)": round(info["MAE (Log)"], 4),
            "RMSE ($)": round(info["RMSE ($)"], 2),
            "MAE ($)": round(info["MAE ($)"], 2),
            "CV RMSE Mean": round(info["CV RMSE Mean"], 4),
            "CV RMSE Std": round(info["CV RMSE Std"], 4),
        }
        for name, info in results.items()
    ]
    _write_dataframe(wb.active, "Model Comparison", pd.DataFrame(summary_data))

    best_row = max(range(2, len(summary_data) + 2), key=lambda row: wb.active.cell(row=row, column=2).value)
    for col_idx in range(1, len(summary_data[0]) + 1):
        wb.active.cell(row=best_row, column=col_idx).fill = best_fill
    style_header(wb.active, len(summary_data[0]))
    auto_width(wb.active)

    hyper_data = []
    for name, model in build_models().items():
        for param, value in model.get_params().items():
            hyper_data.append({"Model": name, "Parameter": param, "Value": str(value)})
    ws = wb.create_sheet("Hyperparameters")
    _write_dataframe(ws, "Hyperparameters", pd.DataFrame(hyper_data))
    style_header(ws, 3)
    auto_width(ws)

    fi_data = []
    for name in ["Random Forest", "XGBoost"]:
        fi_df = pd.DataFrame({
            "Feature": feature_columns,
            "Importance": results[name]["model"].feature_importances_,
        }).sort_values("Importance", ascending=False).head(20)
        for _, row in fi_df.iterrows():
            fi_data.append({
                "Model": name,
                "Feature": row["Feature"],
                "Importance": round(row["Importance"], 6),
            })
    ws = wb.create_sheet("Feature Importance")
    _write_dataframe(ws, "Feature Importance", pd.DataFrame(fi_data))
    style_header(ws, 3)
    auto_width(ws)

    best_name = max(results, key=lambda name: results[name]["R2 Score"])
    dataset_info = pd.DataFrame([
        ["Training Set Size", X_train.shape[0]],
        ["Test Set Size", X_test.shape[0]],
        ["Number of Features", X_train.shape[1]],
        ["Target Variable", "SalePrice (log1p transformed)"],
        ["Test Size Ratio", "20%"],
        ["Random State", 42],
        ["Best Model", best_name],
        ["Best R2 Score", round(results[best_name]["R2 Score"], 4)],
        ["Best RMSE ($)", round(results[best_name]["RMSE ($)"], 2)],
    ], columns=["Property", "Value"])
    ws = wb.create_sheet("Dataset Info")
    _write_dataframe(ws, "Dataset Info", dataset_info)
    style_header(ws, 2)
    auto_width(ws)

    pred_data = {"Actual (Log)": y_test.values, "Actual ($)": np.expm1(y_test).values}
    for name, info in results.items():
        pred_data[f"{name} Pred (Log)"] = info["y_pred"]
        pred_data[f"{name} Pred ($)"] = info["y_pred_orig"]
    ws = wb.create_sheet("Prediction Details")
    _write_dataframe(ws, "Prediction Details", pd.DataFrame(pred_data).round(4))
    style_header(ws, len(pred_data))
    auto_width(ws)

    wb.save(EXCEL_PATH)


def _write_dataframe(ws, title, df):
    ws.title = title
    for col_idx, col_name in enumerate(df.columns, 1):
        ws.cell(row=1, column=col_idx, value=col_name)
    for row_idx, row in enumerate(df.itertuples(index=False), 2):
        for col_idx, value in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)


def run_comparison(show_plots=False, export_excel=True):
    print("=" * 60)
    print("  HOUSE PRICE PREDICTION - MODEL COMPARISON")
    print("=" * 60)

    X, y = prepare_training_data()
    X_train, X_test, y_train, y_test = split_data(X, y)

    print("\n  Dataset Info:")
    print(f"    Total samples  : {len(X)}")
    print(f"    Training set   : {X_train.shape[0]}")
    print(f"    Test set       : {X_test.shape[0]}")
    print(f"    Features       : {X_train.shape[1]}")

    results = evaluate_models(X_train, X_test, y_train, y_test, X, y)

    comparison_df = pd.DataFrame({
        name: {
            "RMSE (Log)": f"{info['RMSE (Log)']:.4f}",
            "R2 Score": f"{info['R2 Score']:.4f}",
            "MAE (Log)": f"{info['MAE (Log)']:.4f}",
            "RMSE ($)": f"${info['RMSE ($)']:,.0f}",
            "MAE ($)": f"${info['MAE ($)']:,.0f}",
            "CV RMSE": f"{info['CV RMSE Mean']:.4f} +/- {info['CV RMSE Std']:.4f}",
        }
        for name, info in results.items()
    }).T

    print(f"\n{'=' * 60}")
    print("  MODEL COMPARISON SUMMARY")
    print(f"{'=' * 60}\n")
    print(comparison_df.to_string())

    best_name = max(results, key=lambda name: results[name]["R2 Score"])
    print(f"\n  >>> Best Model: {best_name}")
    print(f"      R2 Score : {results[best_name]['R2 Score']:.4f}")
    print(f"      RMSE ($) : ${results[best_name]['RMSE ($)']:,.2f}")

    save_plots(results, y_test, X.columns, show_plots=show_plots)
    if export_excel:
        export_results(results, X_train, X_test, y_test, X.columns)

    print(f"\n  Output files:")
    print(f"    Excel  : {EXCEL_PATH}")
    print(f"    Plots  : {EVAL_DIR}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--show-plots", action="store_true")
    parser.add_argument("--no-excel", action="store_true")
    args = parser.parse_args()

    run_comparison(show_plots=args.show_plots, export_excel=not args.no_excel)


if __name__ == "__main__":
    main()
